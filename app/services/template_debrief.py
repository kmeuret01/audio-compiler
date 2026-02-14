"""Full Debrief DOCX Template — editable document with DISCUSSED/DECIDED/NEXT structure."""

from pathlib import Path
from typing import List

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# Brand colors as RGB tuples
COLORS = {
    "black": RGBColor(0x0A, 0x0A, 0x0A),
    "white": RGBColor(0xFA, 0xFA, 0xF9),
    "warm_gray": RGBColor(0xE8, 0xE5, 0xE0),
    "mid_gray": RGBColor(0x9B, 0x95, 0x90),
    "dark_gray": RGBColor(0x2D, 0x29, 0x26),
    "accent": RGBColor(0xC8, 0x5A, 0x2A),
    "accent_light": RGBColor(0xE8, 0x70, 0x4A),
    "green": RGBColor(0x2E, 0x7D, 0x32),
    "blue": RGBColor(0x1E, 0x88, 0xE5),
}


def _set_cell_shading(cell, color_hex: str):
    """Set cell background color."""
    shading = OxmlElement('w:shd')
    shading.set(qn('w:fill'), color_hex.lstrip('#'))
    cell._tc.get_or_add_tcPr().append(shading)


def _add_bottom_border(cell, color_hex: str = "E8E5E0"):
    """Add bottom border to cell."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '4')
    bottom.set(qn('w:color'), color_hex)
    tcBorders.append(bottom)
    tcPr.append(tcBorders)


def generate_debrief(
    output_path: str,
    title: str,
    date: str,
    attendees: List[str],
    duration: str,
    bluf: str,
    topics: List[dict],
    execution_steps: List[dict],
) -> str:
    """
    Generate the Full Debrief DOCX.

    Args:
        output_path: Where to save the DOCX
        title: Meeting title
        date: Meeting date
        attendees: List of attendee names
        duration: Estimated duration string
        bluf: Bottom Line Up Front text (full paragraph)
        topics: List of dicts: {title, discussed, decided, next_steps}
        execution_steps: List of dicts: {num, title, tag, summary}

    Returns:
        Path to generated DOCX
    """
    doc = Document()

    # Set default font
    style = doc.styles['Normal']
    style.font.name = 'DM Sans'
    style.font.size = Pt(11)
    style.font.color.rgb = COLORS["dark_gray"]

    # Logo line: [ ] AI LABS
    logo_para = doc.add_paragraph()
    logo_para.alignment = WD_ALIGN_PARAGRAPH.LEFT

    # Brackets in accent color
    run_brackets = logo_para.add_run("[ ] ")
    run_brackets.font.name = 'Space Mono'
    run_brackets.font.size = Pt(14)
    run_brackets.font.color.rgb = COLORS["accent"]
    run_brackets.bold = True

    # AI in black
    run_ai = logo_para.add_run("AI")
    run_ai.font.name = 'Space Mono'
    run_ai.font.size = Pt(12)
    run_ai.font.color.rgb = COLORS["black"]
    run_ai.bold = True

    # LABS in gray
    run_labs = logo_para.add_run("LABS")
    run_labs.font.name = 'DM Sans'
    run_labs.font.size = Pt(9)
    run_labs.font.color.rgb = COLORS["mid_gray"]

    # Doc type label
    run_type = logo_para.add_run("    FULL DEBRIEF")
    run_type.font.name = 'Space Mono'
    run_type.font.size = Pt(8)
    run_type.font.color.rgb = COLORS["mid_gray"]

    # Title
    title_para = doc.add_paragraph()
    title_run = title_para.add_run(title)
    title_run.font.name = 'DM Sans'
    title_run.font.size = Pt(18)
    title_run.font.color.rgb = COLORS["black"]
    title_run.bold = True

    # Meta line
    attendees_str = ", ".join(attendees[:5])
    if len(attendees) > 5:
        attendees_str += f" +{len(attendees) - 5} more"
    meta_para = doc.add_paragraph()
    meta_run = meta_para.add_run(f"{date}  |  {duration}  |  {attendees_str}")
    meta_run.font.name = 'Space Mono'
    meta_run.font.size = Pt(9)
    meta_run.font.color.rgb = COLORS["mid_gray"]

    # Orange divider (using a thin table)
    divider = doc.add_table(rows=1, cols=1)
    divider.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = divider.rows[0].cells[0]
    cell.width = Inches(6.5)
    _set_cell_shading(cell, "C85A2A")
    cell.paragraphs[0].clear()
    # Set row height to minimal
    row = divider.rows[0]
    row.height = Pt(3)

    doc.add_paragraph()  # Spacer

    # BOTTOM LINE UP FRONT
    bluf_header = doc.add_paragraph()
    bluf_header_run = bluf_header.add_run("BOTTOM LINE UP FRONT")
    bluf_header_run.font.name = 'Space Mono'
    bluf_header_run.font.size = Pt(9)
    bluf_header_run.font.color.rgb = COLORS["accent"]
    bluf_header_run.bold = True

    bluf_para = doc.add_paragraph()
    bluf_run = bluf_para.add_run(bluf)
    bluf_run.font.size = Pt(11)
    bluf_run.font.color.rgb = COLORS["dark_gray"]

    doc.add_paragraph()  # Spacer

    # DISCUSSION TOPICS
    if topics:
        topics_header = doc.add_paragraph()
        topics_header_run = topics_header.add_run("DISCUSSION TOPICS")
        topics_header_run.font.name = 'Space Mono'
        topics_header_run.font.size = Pt(9)
        topics_header_run.font.color.rgb = COLORS["accent"]
        topics_header_run.bold = True

        for topic in topics:
            # Topic title
            topic_title = doc.add_paragraph()
            topic_title_run = topic_title.add_run(topic.get('title', 'Topic'))
            topic_title_run.font.size = Pt(12)
            topic_title_run.font.color.rgb = COLORS["black"]
            topic_title_run.bold = True

            # Create table for DISCUSSED/DECIDED/NEXT
            topic_table = doc.add_table(rows=3, cols=2)
            topic_table.alignment = WD_TABLE_ALIGNMENT.LEFT

            # Set column widths
            topic_table.columns[0].width = Inches(0.9)
            topic_table.columns[1].width = Inches(5.6)

            # DISCUSSED row
            label_cell = topic_table.rows[0].cells[0]
            label_para = label_cell.paragraphs[0]
            label_run = label_para.add_run("DISCUSSED")
            label_run.font.name = 'Space Mono'
            label_run.font.size = Pt(8)
            label_run.font.color.rgb = COLORS["accent"]
            label_run.bold = True

            content_cell = topic_table.rows[0].cells[1]
            content_para = content_cell.paragraphs[0]
            content_run = content_para.add_run(topic.get('discussed', ''))
            content_run.font.size = Pt(10)
            content_run.font.color.rgb = COLORS["dark_gray"]
            _add_bottom_border(content_cell)

            # DECIDED row
            label_cell = topic_table.rows[1].cells[0]
            label_para = label_cell.paragraphs[0]
            label_run = label_para.add_run("DECIDED")
            label_run.font.name = 'Space Mono'
            label_run.font.size = Pt(8)
            label_run.font.color.rgb = COLORS["green"]
            label_run.bold = True

            content_cell = topic_table.rows[1].cells[1]
            content_para = content_cell.paragraphs[0]
            content_run = content_para.add_run(topic.get('decided', ''))
            content_run.font.size = Pt(10)
            content_run.font.color.rgb = COLORS["dark_gray"]
            _add_bottom_border(content_cell)

            # NEXT row
            label_cell = topic_table.rows[2].cells[0]
            label_para = label_cell.paragraphs[0]
            label_run = label_para.add_run("NEXT")
            label_run.font.name = 'Space Mono'
            label_run.font.size = Pt(8)
            label_run.font.color.rgb = COLORS["blue"]
            label_run.bold = True

            content_cell = topic_table.rows[2].cells[1]
            content_para = content_cell.paragraphs[0]
            content_run = content_para.add_run(topic.get('next_steps', ''))
            content_run.font.size = Pt(10)
            content_run.font.color.rgb = COLORS["dark_gray"]
            _add_bottom_border(content_cell)

            doc.add_paragraph()  # Spacer between topics

    # EXECUTION BOARD
    if execution_steps:
        exec_header = doc.add_paragraph()
        exec_header_run = exec_header.add_run("EXECUTION BOARD")
        exec_header_run.font.name = 'Space Mono'
        exec_header_run.font.size = Pt(9)
        exec_header_run.font.color.rgb = COLORS["accent"]
        exec_header_run.bold = True

        for step in execution_steps:
            step_para = doc.add_paragraph()
            step_num = step.get('num', '0')
            step_title = step.get('title', '')
            step_tag = step.get('tag', '')
            step_summary = step.get('summary', '')

            # Step number and title
            num_run = step_para.add_run(f"[{step_num}] ")
            num_run.font.name = 'Space Mono'
            num_run.font.size = Pt(10)
            num_run.font.color.rgb = COLORS["accent"] if step_num == '0' else COLORS["dark_gray"]
            num_run.bold = True

            title_run = step_para.add_run(step_title)
            title_run.font.size = Pt(10)
            title_run.font.color.rgb = COLORS["black"]
            title_run.bold = True

            if step_tag:
                tag_run = step_para.add_run(f"  [{step_tag}]")
                tag_run.font.name = 'Space Mono'
                tag_run.font.size = Pt(8)
                tag_run.font.color.rgb = COLORS["mid_gray"]

            if step_summary:
                summary_para = doc.add_paragraph()
                summary_para.paragraph_format.left_indent = Inches(0.3)
                summary_run = summary_para.add_run(step_summary)
                summary_run.font.size = Pt(10)
                summary_run.font.color.rgb = COLORS["dark_gray"]

    # Footer
    doc.add_paragraph()
    footer_para = doc.add_paragraph()
    footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer_para.add_run("goailabs.ai  |  CONFIDENTIAL")
    footer_run.font.name = 'Space Mono'
    footer_run.font.size = Pt(8)
    footer_run.font.color.rgb = COLORS["mid_gray"]

    doc.save(output_path)
    return output_path
